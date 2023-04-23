import os
from pathlib import Path
from zipfile import ZipFile
from docx import Document
import xlsxwriter


def create_test_dir(TEST_DIRECTORY: Path):
    os.mkdir(TEST_DIRECTORY)
    file_1 = TEST_DIRECTORY / "test1.txt"
    file_2 = TEST_DIRECTORY / "test2.txt"
    with open(file_1, "w+", encoding="utf-8") as file:
        file.write("abracadabra")
    with open(file_2, "w+", encoding="utf-8") as file:
        file.write("abr")
    
    second_folder = TEST_DIRECTORY / "second_folder"
    os.mkdir(second_folder)
    file_3 = second_folder / "test3.txt"
    with open(file_3, "w+", encoding="utf-8") as file:
        file.write("a \n ab \n abr \n abra")
    
    zip_file = TEST_DIRECTORY / "archive.zip"
    with ZipFile(zip_file, 'w') as zip_object:
        zip_object.write(file_1)
        zip_object.write(file_2)
        zip_object.write(file_3)

    document = Document()
    document.add_paragraph("abracadabra")
    document.save(TEST_DIRECTORY / "word_doc.docx")

    workbook = xlsxwriter.Workbook(TEST_DIRECTORY / 'excel_file.xlsx')
    worksheet = workbook.add_worksheet()
    worksheet.write('A1', 'abracadabra')
    workbook.close()
